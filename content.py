import os

from pypdf import PdfReader
from pptx import Presentation


DATA_FOLDER = "./data"
OUTPUT_FILE = "entire_data_content.txt"


def extract_pdf_text(file_path: str) -> str:
    """Extract selectable text from a PDF, page by page."""
    reader = PdfReader(file_path)
    text_blocks = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()

        if text:
            text_blocks.append(f"--- Page {page_number} ---")
            text_blocks.append(text)

    return "\n".join(text_blocks)


def extract_ppt_text(file_path: str) -> str:
    """Extract text and tables from a PPTX presentation."""
    presentation = Presentation(file_path)
    text_blocks = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        slide_blocks = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()

                    if line:
                        slide_blocks.append(line)

            if shape.has_table:
                table_lines = []

                for row in shape.table.rows:
                    row_data = [
                        cell.text.strip()
                        for cell in row.cells
                    ]

                    if any(row_data):
                        table_lines.append(" | ".join(row_data))

                if table_lines:
                    slide_blocks.append("[Table]:")
                    slide_blocks.extend(table_lines)

        if slide_blocks:
            text_blocks.append(f"--- Slide {slide_number} ---")
            text_blocks.extend(slide_blocks)

    return "\n".join(text_blocks)


def process_data_folder():
    """Extract text from PDF and PPTX files in the data folder."""
    if not os.path.exists(DATA_FOLDER):
        os.makedirs(DATA_FOLDER, exist_ok=True)
        print(
            f"Created '{DATA_FOLDER}'. "
            "Place your PDF and PPTX files inside it."
        )
        return

    files = [
        filename
        for filename in sorted(os.listdir(DATA_FOLDER))
        if not filename.startswith(".")
        and os.path.isfile(os.path.join(DATA_FOLDER, filename))
    ]

    if not files:
        print(f"No files found inside '{DATA_FOLDER}'.")
        return

    all_extracted_text = []

    for filename in files:
        file_path = os.path.join(DATA_FOLDER, filename)
        file_ext = os.path.splitext(filename)[1].lower()

        if file_ext not in {".pdf", ".pptx"}:
            print(f"Skipping unsupported file: {filename}")
            continue

        print(f"Processing: {filename}...")

        try:
            if file_ext == ".pdf":
                content = extract_pdf_text(file_path)
            else:
                content = extract_ppt_text(file_path)

            if not content.strip():
                print(f"No readable text found: {filename}")
                continue

            all_extracted_text.extend(
                [
                    "=" * 40,
                    f"FILE: {filename}",
                    "=" * 40,
                    content,
                    "",
                ]
            )

        except Exception as exc:
            print(f"Failed to process {filename}: {exc}")

    if not all_extracted_text:
        print("No readable content was extracted.")
        return

    with open(OUTPUT_FILE, "w", encoding="utf-8") as output_file:
        output_file.write("\n".join(all_extracted_text))

    print(
        "\nDone! Extracted text saved to: "
        f"{os.path.abspath(OUTPUT_FILE)}"
    )


if __name__ == "__main__":
    process_data_folder()